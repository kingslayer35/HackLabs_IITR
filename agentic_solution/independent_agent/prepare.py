from __future__ import annotations

import csv
import hashlib
import json
import random
import re
import sqlite3
import time
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .common import ARTIFACTS, ROOT, SEED, SYSTEMS, digest, input_id, jaccard, norm, shingles, write_json

TAG_NAMES = set("""disorder|finding|procedure|situation|observable entity|body structure|morphologic abnormality|qualifier value|substance|organism|physical object|regime/therapy|event|occupation|environment|assessment scale|navigational concept|record artifact|person|physical force|attribute|clinical drug|product|cell|cell structure|specimen|staging scale|tumor staging|link assertion|metadata|foundation metadata concept|special concept|namespace concept|administrative concept|medicinal product|medicinal product form|clinical finding""".split("|"))


def tag_of(description: str):
    match = re.search(r"\(([^()]+)\)$", description)
    return match.group(1) if match and match.group(1) in TAG_NAMES else None


def qualifier_family(description: str) -> str:
    description = re.sub(r"\s*\([^()]+\)$", "", description)
    text = norm(description)
    text = re.sub(r"\b(family history of|history of|left|right|bilateral|acute|chronic|absent|absence of|no)\b", " ", text)
    return " ".join(text.split())


def build_splits(data_dir: Path, out: Path, count: int = 20) -> dict:
    with (data_dir / "hackathon_train.csv").open(newline="") as handle:
        source = list(csv.DictReader(handle))
    rows = []
    for i, row in enumerate(source):
        if row["subconcept"] != "DiagnosisCode":
            continue
        row = {**row, "source_row": i + 2}
        row["input_id"] = input_id(row["entity"], row["chunk_text"])
        # Blank contexts are not treated as a shared document.
        row["group_id"] = digest(norm(row["chunk_text"]) or "empty:" + norm(row["entity"]))[:20]
        rows.append(row)
    group_ids = sorted({r["group_id"] for r in rows})
    random.Random(SEED).shuffle(group_ids)
    n = len(group_ids)
    assignments = {g: "train" if i < int(.70 * n) else "dev" if i < int(.85 * n) else "test"
                   for i, g in enumerate(group_ids)}
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["input_id"]].append(row)
    eligible = [key for key, rr in grouped.items()
                if assignments[rr[0]["group_id"]] == "test" and {r["coding_system"] for r in rr} == set(SYSTEMS)]
    eligible.sort(key=lambda key: digest(f"smoke:{SEED}:{key}"))
    if len(eligible) < count:
        raise ValueError("Not enough paired held-out inputs")
    chosen_ids = set(eligible[:count])
    chosen_rows = [r for r in rows if r["input_id"] in chosen_ids]
    chosen_entities = {norm(r["entity"]) for r in chosen_rows}
    chosen_shingles = {r["input_id"]: shingles(r["chunk_text"]) for r in chosen_rows}
    # Audit ALL candidate example chunks against ALL 20 test inputs. This is an
    # exhaustive check, rather than an approximate duplicate search.
    group_shingles = {r["group_id"]: shingles(r["chunk_text"]) for r in rows}
    excluded_groups = set()
    duplicate_matches = []
    for group, candidate in group_shingles.items():
        if assignments[group] != "train":
            continue
        for key, target in chosen_shingles.items():
            score = jaccard(candidate, target)
            if score >= .85:
                excluded_groups.add(group)
                duplicate_matches.append({"group_id": group, "test_id": key, "jaccard": score})
                break
    # Also exclude any whole training group containing an identical normalized
    # mention, even if it came from a different chart or coding system.
    for row in rows:
        if assignments[row["group_id"]] == "train" and norm(row["entity"]) in chosen_entities:
            excluded_groups.add(row["group_id"])
    examples = [r for r in rows if assignments[r["group_id"]] == "train" and r["group_id"] not in excluded_groups]
    smoke = []
    for key in eligible[:count]:
        rr = grouped[key]
        smoke.append({"input_id": key, "entity": rr[0]["entity"], "chunk_text": rr[0]["chunk_text"],
                      "group_id": rr[0]["group_id"], "gold": {r["coding_system"]: r["code"] for r in rr}})
    example_groups = {r["group_id"] for r in examples}
    test_groups = {r["group_id"] for r in chosen_rows}
    example_entities = {norm(r["entity"]) for r in examples}
    max_similarity = max((jaccard(group_shingles[g], target) for g in example_groups
                          for target in chosen_shingles.values()), default=0.0)
    audit = {
        "seed": SEED, "smoke_inputs": len(smoke), "selection": "seeded hash order over paired inputs in the test split; no outcome-based selection",
        "diagnosis_rows": len(rows), "source_rows_by_system": dict(Counter(r["coding_system"] for r in rows)),
        "source_group_counts": dict(Counter(assignments.values())), "training_example_rows": len(examples),
        "quarantined_training_groups": len(excluded_groups), "near_duplicate_training_groups": len(duplicate_matches),
        "shared_input_ids": len(chosen_ids & {r["input_id"] for r in examples}),
        "shared_normalized_mentions": len(chosen_entities & example_entities),
        "shared_normalized_context_groups": len(example_groups & test_groups),
        "near_duplicate_threshold": .85, "maximum_test_training_shingle_jaccard": max_similarity,
        "near_duplicate_check": "exhaustive test-to-training 3-word-shingle Jaccard; groups at >= 0.85 quarantined",
        "scope": "leakage audit covers the 20-input smoke test; other dev/test rows are reserved and not audited for near duplicates here",
    }
    assert audit["shared_input_ids"] == audit["shared_normalized_mentions"] == audit["shared_normalized_context_groups"] == 0
    assert max_similarity < .85
    out.mkdir(parents=True, exist_ok=True)
    write_json(out / "smoke20.json", smoke)
    write_json(out / "training_examples.json", examples)
    write_json(out / "split_manifest.json", {"seed": SEED, "groups": assignments, "quarantined_training_groups": sorted(excluded_groups),
                                              "rows": [{k: r[k] for k in ("input_id", "group_id", "source_row")} for r in rows]})
    write_json(out / "leakage_audit.json", audit)
    return audit


def build_terminology(data_dir: Path, out: Path) -> dict:
    path = out / "terminology.sqlite"
    if path.exists():
        raise FileExistsError(f"Remove {path} explicitly to rebuild terminology")
    con = sqlite3.connect(path)
    con.executescript("""
        PRAGMA journal_mode=OFF; PRAGMA synchronous=OFF; PRAGMA temp_store=MEMORY;
        CREATE TABLE concepts(system TEXT, code TEXT, preferred TEXT, fsn TEXT, tag TEXT, parent TEXT, is_leaf INTEGER, family TEXT, PRIMARY KEY(system,code));
        CREATE TABLE aliases(id INTEGER PRIMARY KEY, system TEXT, code TEXT, text TEXT, normalized TEXT, description_order INTEGER);
        CREATE TABLE edges(system TEXT, source TEXT, relation TEXT, target TEXT, provenance TEXT, PRIMARY KEY(system,source,relation,target));
        CREATE TABLE examples(id INTEGER PRIMARY KEY, system TEXT, code TEXT, entity TEXT, context TEXT, input_id TEXT, group_id TEXT);
        CREATE VIRTUAL TABLE alias_fts USING fts5(text, content='aliases', content_rowid='id', tokenize='porter unicode61 remove_diacritics 2');
        CREATE VIRTUAL TABLE example_fts USING fts5(entity, context, content='examples', content_rowid='id', tokenize='porter unicode61 remove_diacritics 2');
    """)
    started = time.monotonic()
    stats = {}
    for system, filename in [("ICD10CM", "icd10cm_codes.xlsx"), ("SNOMEDCT_US", "SNOMEDCT_US_codes_and_descriptions.xlsx")]:
        print(f"Loading {system} dictionary", flush=True)
        workbook = load_workbook(data_dir / filename, read_only=True, data_only=True)
        concepts = {}
        batch = []
        total = 0
        for values in workbook.worksheets[0].iter_rows(min_row=2, values_only=True):
            if values[0] is None or values[1] is None:
                continue
            code, description = str(values[0]).strip(), str(values[1]).strip()
            order = int(values[2]) if system == "SNOMEDCT_US" else 1
            tag = tag_of(description) if system == "SNOMEDCT_US" else None
            if code not in concepts:
                concepts[code] = [description, description if tag else None, tag]
            elif tag:
                concepts[code][1:] = [description, tag]
            batch.append((system, code, description, norm(description), order))
            total += 1
            if len(batch) >= 10000:
                con.executemany("INSERT INTO aliases(system,code,text,normalized,description_order) VALUES(?,?,?,?,?)", batch)
                batch.clear()
                if total % 100000 == 0:
                    print(f"  {total:,} descriptions", flush=True)
        if batch:
            con.executemany("INSERT INTO aliases(system,code,text,normalized,description_order) VALUES(?,?,?,?,?)", batch)
        workbook.close()
        parents = {}
        if system == "ICD10CM":
            compact_to_code = {c.replace(".", ""): c for c in concepts if "-" not in c}
            for code in concepts:
                if "-" in code:
                    continue
                compact = code.replace(".", "")
                for n in range(len(compact) - 1, 2, -1):
                    parent = compact_to_code.get(compact[:n])
                    if parent:
                        parents[code] = parent
                        break
            ranges = []
            for code in concepts:
                if "-" in code:
                    left, right = code.split("-", 1)
                    ranges.append((left[:3], right[:3], code))
            for code in concepts:
                if len(code) == 3:
                    containing = [(lo, hi, key) for lo, hi, key in ranges if lo <= code <= hi]
                    if containing:
                        parents[code] = min(containing, key=lambda item: (item[1], item[0]))[2]
            con.executemany("INSERT INTO edges VALUES(?,?,?,?,?)",
                            [(system, child, "is_a", parent, "derived ICD code prefix or supplied chapter range") for child, parent in sorted(parents.items())])
        nonleaves = set(parents.values())
        con.executemany("INSERT INTO concepts VALUES(?,?,?,?,?,?,?,?)", [
            (system, code, values[0], values[1], values[2], parents.get(code),
             int(code not in nonleaves and "-" not in code) if system == "ICD10CM" else None,
             qualifier_family(values[0]) if system == "SNOMEDCT_US" else None)
            for code, values in concepts.items()])
        stats[system] = {"concepts": len(concepts), "descriptions": total, "hierarchy_edges": len(parents),
                         "semantic_tags": sum(v[2] is not None for v in concepts.values())}
        con.commit()
    print("Building search and relationship indexes", flush=True)
    con.executescript("""
        CREATE INDEX aliases_concept ON aliases(system,code);
        CREATE INDEX aliases_exact ON aliases(system,normalized);
        CREATE INDEX concept_parent ON concepts(system,parent);
        CREATE INDEX concept_family ON concepts(system,family,tag);
        INSERT INTO alias_fts(alias_fts) VALUES('rebuild');
        INSERT INTO alias_fts(alias_fts) VALUES('optimize');
    """)
    examples = json.loads((out / "training_examples.json").read_text())
    valid = []
    for example in examples:
        if con.execute("SELECT 1 FROM concepts WHERE system=? AND code=?", (example["coding_system"], example["code"])).fetchone():
            valid.append((example["coding_system"], example["code"], example["entity"], example["chunk_text"], example["input_id"], example["group_id"]))
    con.executemany("INSERT INTO examples(system,code,entity,context,input_id,group_id) VALUES(?,?,?,?,?,?)", valid)
    con.execute("INSERT INTO example_fts(example_fts) VALUES('rebuild')")
    con.commit()
    stats["valid_training_examples"] = len(valid)
    stats["excluded_invalid_training_examples"] = len(examples) - len(valid)
    stats["build_seconds"] = round(time.monotonic() - started, 2)
    stats["sources"] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(data_dir.glob("*")) if p.is_file()}
    stats["relationship_limits"] = "ICD hierarchy derived from supplied code structure. No official SNOMED relationship edges supplied; qualifier families are derived search hints only."
    write_json(out / "terminology_stats.json", stats)
    con.close()
    print(json.dumps(stats, indent=2), flush=True)
    return stats


def prepare(data_dir: Path = ROOT / "data", out: Path = ARTIFACTS, count: int = 20):
    audit = build_splits(data_dir, out, count)
    print(json.dumps(audit, indent=2), flush=True)
    return build_terminology(data_dir, out)
